import os
import json
import io
import logging
from flask import Flask, render_template, request, session, send_file, redirect, url_for
from flask_session import Session
from flask_limiter import Limiter
from flask_limiter.util import get_remote_address
from openai import OpenAI
from dotenv import load_dotenv
import pypdf
from docx import Document
from fpdf import FPDF

load_dotenv()
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

app = Flask(__name__)
app.secret_key = os.getenv("SECRET_KEY", "anas_secret_2026")
app.config["MAX_CONTENT_LENGTH"] = 5 * 1024 * 1024
app.config["SESSION_TYPE"] = "cachelib"
app.config["SESSION_PERMANENT"] = False
Session(app)

limiter = Limiter(
    get_remote_address,
    app=app,
    default_limits=["60 per minute"],
    storage_uri="memory://"
)

client = OpenAI(
    base_url="https://api.groq.com/openai/v1",
    api_key=os.getenv("GROQ_API_KEY")
)

ALLOWED_EXT = {"pdf", "docx", "txt", "md"}

def get_extension(filename: str) -> str:
    parts = filename.rsplit(".", 1)
    return parts[1].lower() if len(parts) == 2 else ""

def extract_text(file) -> str:
    if not file or file.filename == "":
        return ""
    ext = get_extension(file.filename)
    if ext not in ALLOWED_EXT:
        return ""
    try:
        if ext == "pdf":
            reader = pypdf.PdfReader(file)
            return "\n".join(p.extract_text() for p in reader.pages if p.extract_text())
        elif ext == "docx":
            doc = Document(file)
            return "\n".join(p.text for p in doc.paragraphs if p.text)
        elif ext in ("txt", "md"):
            return file.read().decode("utf-8", errors="replace")
    except Exception as e:
        logger.error(f"Extract error: {e}")
    return ""

def analyze_cv(cv_text: str, job_desc: str = "") -> dict | None:
    system = """You are a Finnish Career Expert. Analyze the CV.
    Return JSON ONLY:
    {
        "score": int,
        "summary": "str",
        "strengths": ["exactly 4 items"],
        "weaknesses": ["exactly 4 items"],
        "improvements": ["exactly 5 actions"],
        "recommended_roles": [
            {
                "role": "Job Title 1",
                "companies": ["exactly 25 unique Finnish company names, NO duplicates, NO URLs"]
            },
            {
                "role": "Job Title 2",
                "companies": ["exactly 25 unique Finnish company names, NO duplicates, NO URLs"]
            },
            {
                "role": "Job Title 3",
                "companies": ["exactly 25 unique Finnish company names, NO duplicates, NO URLs"]
            },
            {
                "role": "Job Title 4",
                "companies": ["exactly 25 unique Finnish company names, NO duplicates, NO URLs"]
            }
        ]
    }

    IMPORTANT:
    - Return EXACTLY 4 roles.
    - Each role must have EXACTLY 25 company names.
    - Company names only, NO URLs, NO websites, NO links.
    - NO duplicate companies within the same role.
    - Use only realistic Finnish employers (e.g. KONE, Nokia, Elisa, Fortum, Wartsila, OP Group,
      UPM, Metso, Konecranes, YIT, Stora Enso, Neste, Valmet, Nordea, Tietoevry, Reaktor,
      Futurice, Wolt, Supercell, Rovio, ABB Finland, Outokumpu, Kemira, Huhtamaki, Orion, etc.).
    - Output must be valid JSON only."""

    try:
        res = client.chat.completions.create(
            model="llama-3.3-70b-versatile",
            messages=[
                {"role": "system", "content": system},
                {"role": "user", "content": f"CV:\n{cv_text}\n\nJob Description:\n{job_desc}"}
            ],
            temperature=0.2,
            response_format={"type": "json_object"},
            max_tokens=6000
        )
        result = json.loads(res.choices[0].message.content)
        required = ["score", "summary", "strengths", "weaknesses", "improvements", "recommended_roles"]
        if not all(k in result for k in required):
            logger.error("AI response missing keys")
            return None
        result["score"] = max(0, min(100, int(result.get("score", 0))))

        for role in result.get("recommended_roles", []):
            seen = []
            for c in role.get("companies", []):
                if c not in seen:
                    seen.append(c)
            role["companies"] = seen[:25]

        return result
    except json.JSONDecodeError as e:
        logger.error(f"JSON parse error: {e}")
    except Exception as e:
        logger.error(f"AI error: {e}")
    return None

@app.route("/", methods=["GET", "POST"])
@limiter.limit("20 per minute")
def home():
    feedback = session.get("feedback")
    cv_text  = session.get("cv_text", "")
    job_desc = session.get("job_desc", "")
    error    = None

    if request.method == "POST":
        cv_f = request.files.get("cv_file")
        jd_f = request.files.get("job_file")
        cv_text  = extract_text(cv_f) if (cv_f and cv_f.filename) else request.form.get("cv_text", "")
        job_desc = extract_text(jd_f) if (jd_f and jd_f.filename) else request.form.get("job_desc", "")

        if cv_text.strip():
            result = analyze_cv(cv_text, job_desc)
            if result:
                session["feedback"] = result
                session["cv_text"]  = cv_text
                session["job_desc"] = job_desc
                return redirect(url_for("home"))
            else:
                error = "Analysis failed. Please try again."
        else:
            error = "Please provide CV content."

    return render_template("index.html", feedback=feedback, cv_text=cv_text, job_desc=job_desc, error=error)

@app.route("/reset")
def reset():
    session.clear()
    return redirect(url_for("home"))

@app.errorhandler(413)
def too_large(e):
    return render_template("index.html",
                           error="File too large. Max 5MB.",
                           feedback=None, cv_text="", job_desc=""), 413

@app.errorhandler(429)
def rate_limited(e):
    return render_template("index.html",
                           error="Too many requests. Please wait.",
                           feedback=None, cv_text="", job_desc=""), 429

FONT_DIR = os.path.join(os.path.dirname(__file__), "static", "fonts")

class MetropoliaPDF(FPDF):
    def __init__(self):
        super().__init__()
        f_reg  = os.path.join(FONT_DIR, "DejaVuSans.ttf")
        f_bold = os.path.join(FONT_DIR, "DejaVuSans-Bold.ttf")
        f_ital = os.path.join(FONT_DIR, "DejaVuSans-Oblique.ttf")
        self.uni = all(os.path.exists(p) for p in [f_reg, f_bold, f_ital])
        if self.uni:
            self.add_font("DejaVu", "",  f_reg,  uni=True)
            self.add_font("DejaVu", "B", f_bold, uni=True)
            self.add_font("DejaVu", "I", f_ital, uni=True)
            self.ff = "DejaVu"
        else:
            self.ff = "Helvetica"

    def use_font(self, style: str = "", size: int = 10):
        self.set_font(self.ff, style, size)

    def header(self):
        self.set_fill_color(255, 80, 0)
        self.rect(0, 0, 210, 5, "F")
        self.set_y(10)
        self.use_font("B", 16)
        self.set_text_color(0, 0, 0)
        self.cell(0, 10, "Metropolia Career AI Analysis Report", ln=True, align="C")

    def footer(self):
        self.set_y(-15)
        self.use_font("I", 8)
        self.set_text_color(150)
        self.cell(0, 10, "Built by Anas Abdelrahim for Metropolia UAS - 2026", align="C")

def draw_role_column(pdf, role_data, x, y, col_width):
    pdf.set_xy(x, y)
    pdf.use_font("B", 12)
    pdf.set_text_color(255, 80, 0)
    pdf.cell(col_width, 8, role_data.get("role", ""), ln=True)
    pdf.set_x(x)
    pdf.use_font("", 8)
    pdf.set_text_color(30, 30, 30)
    for i, company in enumerate(role_data.get("companies", [])[:25], 1):
        pdf.set_x(x)
        pdf.cell(col_width, 5, f"{i:2d}. {company}", ln=True)

@app.route("/download")
def download():
    data = session.get("feedback")
    if not data:
        return "No analysis found.", 400

    pdf = MetropoliaPDF()
    pdf.set_margins(12, 15, 12)
    pdf.set_auto_page_break(False)
    pdf.add_page()

    pdf.use_font("B", 18)
    pdf.set_text_color(255, 80, 0)
    pdf.cell(0, 12, f"Overall Score: {data.get('score')}%", ln=True, align="C")
    pdf.ln(4)

    pdf.use_font("", 11)
    pdf.set_text_color(30, 30, 30)
    pdf.multi_cell(0, 7, txt=str(data.get("summary", "")), align="C")
    pdf.ln(6)

    def add_sec(title, key, bullet, color):
        items = data.get(key, [])
        if not items:
            return
        pdf.use_font("B", 13)
        pdf.set_text_color(*color)
        pdf.cell(0, 9, title, ln=True)
        pdf.ln(1)
        pdf.use_font("", 10)
        pdf.set_text_color(50, 50, 50)
        for item in items:
            pdf.cell(8, 6, bullet, ln=0)
            txt_w = pdf.w - pdf.l_margin - pdf.r_margin - 8
            pdf.multi_cell(txt_w, 6, txt=str(item))
            pdf.ln(1)
        pdf.ln(4)

    add_sec("Key Strengths",         "strengths",    "✓", (25, 135, 84))
    add_sec("Areas for Improvement", "weaknesses",   "△", (220, 53, 69))
    add_sec("Recommended Actions",   "improvements", "→", (255, 80, 0))

    roles   = data.get("recommended_roles", [])
    col_w   = 88
    col1_x  = 12
    col2_x  = 110
    start_y = 35

    if roles:
        pdf.add_page()
        pdf.use_font("B", 14)
        pdf.set_text_color(255, 80, 0)
        pdf.cell(0, 10, "Finnish Market Guide — Suggested Employers", ln=True, align="C")
        pdf.ln(3)
        if len(roles) > 0:
            draw_role_column(pdf, roles[0], col1_x, start_y, col_w)
        if len(roles) > 1:
            draw_role_column(pdf, roles[1], col2_x, start_y, col_w)

        pdf.add_page()
        pdf.use_font("B", 14)
        pdf.set_text_color(255, 80, 0)
        pdf.cell(0, 10, "Finnish Market Guide — Suggested Employers", ln=True, align="C")
        pdf.ln(3)
        if len(roles) > 2:
            draw_role_column(pdf, roles[2], col1_x, start_y, col_w)
        if len(roles) > 3:
            draw_role_column(pdf, roles[3], col2_x, start_y, col_w)

    buf = io.BytesIO(pdf.output())
    return send_file(buf, as_attachment=True,
                     download_name="Metropolia_Career_Report.pdf",
                     mimetype="application/pdf")

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port, debug=False)
